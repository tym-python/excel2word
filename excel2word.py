# !/usr/bin/env python3
# -*-coding:utf-8 -*-

"""
# File       : main.py
# Func     ：
"""
import os
import sys
import re
from datetime import datetime
import streamlit as st
from openpyxl import load_workbook
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Mm


class excel2word():
    def __init__(self,
                 DATE_FORMAT='%Y年%m月%d日',
                 input_DIR_NAME='',
                 OUTPUT_DIR_NAME='output',
                 Image_DIR_NAME='img',
                 base_dir=None):
        '''
        :param DATE_FORMAT:      日期格式
        :param input_DIR_NAME:   输入目录。可为相对 base_dir 的名称，也可为绝对路径。
                                 默认 '' → 使用 base_dir 本身。
        :param OUTPUT_DIR_NAME:  输出目录。可为相对 base_dir 的名称，也可为绝对路径。
                                 默认 'output' → base_dir/output。
        :param Image_DIR_NAME:   插图目录。同上。
        :param base_dir:         基准目录。为 None 时由 get_base_dir() 推断。
        '''
        self.DATE_FORMAT = DATE_FORMAT

        # 基准目录：外部可显式指定；否则自动推断
        self.base_dir = os.path.abspath(base_dir) if base_dir else self.get_base_dir()

        # os.path.join 遇到绝对路径会直接返回该绝对路径（POSIX / Windows 都成立）
        # 因此下面三个参数既支持相对名称，也支持绝对路径，无需额外判断。
        self.input_dir = os.path.join(self.base_dir, input_DIR_NAME)
        self.output_dir = os.path.join(self.base_dir, OUTPUT_DIR_NAME)
        self.img_dir = os.path.join(self.base_dir, Image_DIR_NAME)

    # ============ 路径相关 ============
    def get_base_dir(self):
        """
        获取"运行入口文件"所在目录。
        优先级：
          1. 环境变量 E2W_BASE_DIR（外部框架可注入）
          2. PyInstaller 打包后的可执行文件目录
          3. Streamlit 入口脚本目录
          4. 普通 Python 脚本的入口脚本目录（sys.argv[0]）
          5. 兜底：本文件所在目录
        """
        env_dir = os.environ.get('E2W_BASE_DIR')
        if env_dir:
            return os.path.abspath(env_dir)

        if getattr(sys, 'frozen', False):
            return os.path.dirname(os.path.abspath(sys.executable))

        try:
            from streamlit.runtime.scriptrunner import get_script_run_ctx
            ctx = get_script_run_ctx(suppress_warning=True)
            if ctx and getattr(ctx, 'main_script_path', None):
                return os.path.dirname(os.path.abspath(ctx.main_script_path))
        except Exception:
            pass

        argv0 = sys.argv[0] if sys.argv else ''
        if argv0 and os.path.exists(argv0):
            base = os.path.dirname(os.path.abspath(argv0))
            if os.path.basename(argv0).lower() not in ('streamlit', 'streamlit.exe'):
                return base

        return os.path.dirname(os.path.abspath(__file__))

    # ============ 工具方法 ============
    def find_key_containing(self, data, keyword):
        """从字典的键中查找包含 keyword 的键（不区分大小写），返回第一个匹配的键，找不到返回 None"""
        if isinstance(data, dict):
            keyword_lower = keyword.lower()
            for key in data.keys():
                if keyword_lower in key.lower():
                    return key
        elif isinstance(data, list):
            keyword_lower = keyword.lower()
            for key in data:
                if keyword_lower in key.lower():
                    return key
        return None

    def format_cell_value(self, cell, doc=None, date_format=None):
        """处理数据格式：文本、日期、图片"""
        if not date_format:
            date_format = self.DATE_FORMAT
        value = cell.value
        if value is None:
            return ''
        if isinstance(value, datetime):
            result = date_format
            result = result.replace('%Y', str(value.year))
            result = result.replace('%m', f"{value.month}")
            result = result.replace('%d', f"{value.day}")
            return result
        elif re.findall('png$|jpg$|jpeg$', str(value).lower()):
            img_path = self.get_imagePath(self.img_dir, str(value))
            if img_path and doc:
                img = InlineImage(
                    doc,
                    image_descriptor=img_path,
                    width=Mm(50),
                )
                return img
            else:
                return str(value)
        elif isinstance(value, (int, float)):
            return value
        return value

    def get_outline(self, doc):
        """从 Document 对象中提取标题大纲"""
        outline = []
        for para in doc.paragraphs:
            style_name = para.style.name if para.style else ''
            if style_name.startswith('Heading'):
                try:
                    level = int(style_name.split()[-1])
                except Exception:
                    level = 1
                text = para.text.strip()
                if text:
                    outline.append((level, text))
        if outline:
            highest_level = min(level for level, _ in outline)
            first_highest = next(text for level, text in outline if level == highest_level)
            safe_title = "".join(c for c in first_highest if c not in r'\/:*?"<>|')
            return safe_title
        else:
            return ''

    def get_imagePath(self, img_dir, imgName):
        img_path = os.path.join(img_dir, imgName)
        if not os.path.exists(img_path):
            return None
        else:
            return img_path

    # ============ 主流程 ============
    def main(self, excelName='data.xlsx', wordName='template.docx'):

        excel_path = os.path.join(self.input_dir, excelName)
        template_path = os.path.join(self.input_dir, wordName)

        if not os.path.exists(excel_path):
            st.error(f"[ERROR]：找不到Excel文件 {excel_path}")
            return
        if not os.path.exists(template_path):
            st.error(f"[ERROR]：找不到Word模板 {template_path}")
            return

        os.makedirs(self.output_dir, exist_ok=True)
        os.makedirs(self.img_dir, exist_ok=True)

        wb = load_workbook(excel_path, data_only=True)
        ws = wb.active

        headers = [cell.value for cell in ws[1] if cell.value is not None]
        if not headers:
            st.error("[ERROR]：Excel第一行没有列标题")
            return
        # print(f"[OK] 检测到列标题：{headers}")

        rows = list(ws.iter_rows(min_row=2, values_only=False))

        count = 0
        for idx, row in enumerate(rows, start=2):
            if all(cell.value is None or str(cell.value).strip() == '' for cell in row):
                continue

            doc = DocxTemplate(template_path)
            data = {}
            for i, header in enumerate(headers):
                if i < len(row):
                    cell = row[i]
                    data[header] = self.format_cell_value(cell, doc)

            doc.render(data)

            custom_title = data.get(self.find_key_containing(data, '文件标题'), '').strip()
            if custom_title:
                safe_title = "".join(c for c in custom_title if c not in r'\/:*?"<>|')
            else:
                safe_title = self.get_outline(doc)

            if not safe_title:
                if '单位名称' in data.keys():
                    unit = data.get('单位名称', '').strip()
                else:
                    unit_col = self.find_key_containing(data, '单位')
                    unit = data.get(unit_col, '').strip()
                yearName = str(data.get('年份', '')).strip()
                if unit:
                    safe_unit = "".join(c for c in unit + yearName if c not in r'\/:*?"<>|')
                    output_name = f"{safe_unit}.docx"
                else:
                    output_name = f"doc_{idx}.docx"
            else:
                output_name = f"{safe_title}.docx"

            output_path = os.path.join(self.output_dir, output_name)

            counter = 1
            original_path = output_path
            while os.path.exists(output_path):
                name, ext = os.path.splitext(original_path)
                output_path = f"{name}_{counter}{ext}"
                counter += 1

            doc.save(output_path)
            st.info(f"[OK] 已生成：{output_path}")
            count += 1

        st.success( f"[DONE] 全部完成！共生成 {count} 个文档。")
        # print(f"输出目录：{self.output_dir}")
        wb.close()
        st.info(f'----【{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}】运行结束。----',)
        return count

if __name__ == '__main__':
    print(f'----【{datetime.now()}】开始运行。----')
    c = excel2word(input_DIR_NAME='')
    c.main()